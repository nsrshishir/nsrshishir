import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_spm_docx(output_path):
    doc = docx.Document()

    # Standard A4 page size (210mm x 297mm = 8.27 in x 11.69 in)
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.40)
        section.bottom_margin = Inches(0.40)
        section.left_margin = Inches(0.50)
        section.right_margin = Inches(0.50)

    # Color Palette: Executive Navy & Blue
    NAVY = RGBColor(15, 23, 42)       # #0f172a
    BLUE = RGBColor(37, 99, 235)      # #2563eb
    SLATE = RGBColor(71, 85, 105)     # #475569
    DARK = RGBColor(30, 41, 59)       # #1e293b
    MUTED = RGBColor(100, 116, 139)   # #64748b

    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def clear_cell_borders(cell):
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>')
        tcPr.append(tcBorders)

    def set_cell_margins(cell, top=20, bottom=20, left=35, right=35):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    # 1. HEADER SECTION
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False

    cell_left = header_table.cell(0, 0)
    cell_right = header_table.cell(0, 1)
    cell_left.width = Inches(4.70)
    cell_right.width = Inches(2.57)
    clear_cell_borders(cell_left)
    clear_cell_borders(cell_right)
    set_cell_margins(cell_left, top=10, bottom=20, left=0, right=20)
    set_cell_margins(cell_right, top=10, bottom=20, left=20, right=0)

    # Left: Name & Title
    p_name = cell_left.paragraphs[0]
    p_name.paragraph_format.space_after = Pt(1)
    p_name.paragraph_format.space_before = Pt(0)
    r_name = p_name.add_run("Md Niaj Shahriar Shishir")
    r_name.font.name = "Arial"
    r_name.font.size = Pt(17.5)
    r_name.font.bold = True
    r_name.font.color.rgb = NAVY

    p_title = cell_left.add_paragraph()
    p_title.paragraph_format.space_after = Pt(2)
    p_title.paragraph_format.space_before = Pt(0)
    r_title = p_title.add_run("Senior Project Manager & Technology Lead")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(9.5)
    r_title.font.bold = True
    r_title.font.color.rgb = BLUE

    p_cred = cell_left.add_paragraph()
    p_cred.paragraph_format.space_after = Pt(0)
    p_cred.paragraph_format.space_before = Pt(0)
    r_cred = p_cred.add_run("Certified Odoo Consultant | Software Delivery, Architecture & Team Leadership | 9+ Years Experience")
    r_cred.font.name = "Calibri"
    r_cred.font.size = Pt(8.2)
    r_cred.font.color.rgb = SLATE

    # Right: Contact Info
    contact_lines = [
        ("Portfolio: ", "nsrshishir.github.io/nsrshishir", True),
        ("Email: ", "niaj@metamorphosis.com.bd", False),
        ("Alt Email: ", "nsrshishir550@gmail.com", False),
        ("Phone: ", "+880 1827-512928 | Dhaka, BD", False),
        ("LinkedIn: ", "linkedin.com/in/nsrshishir", False),
        ("GitHub: ", "github.com/nsrshishir", False),
    ]

    p_contact = cell_right.paragraphs[0]
    for idx, (label, val, highlight) in enumerate(contact_lines):
        if idx > 0:
            p_contact = cell_right.add_paragraph()
        p_contact.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_contact.paragraph_format.space_after = Pt(0.5)
        p_contact.paragraph_format.space_before = Pt(0)

        rl = p_contact.add_run(label)
        rl.font.name = "Calibri"
        rl.font.size = Pt(7.6)
        rl.font.bold = True
        rl.font.color.rgb = MUTED

        rv = p_contact.add_run(val)
        rv.font.name = "Calibri"
        rv.font.size = Pt(7.6)
        if highlight:
            rv.font.bold = True
            rv.font.color.rgb = BLUE
        else:
            rv.font.color.rgb = DARK

    # Top Divider Line
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(2)
    p_div.paragraph_format.space_after = Pt(4)
    p_div_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="10" w:space="1" w:color="2563EB"/></w:pBdr>')
    p_div._p.get_or_add_pPr().append(p_div_border)

    # 2. KEY METRICS CALLOUT STRIP
    stats_table = doc.add_table(rows=1, cols=4)
    stats_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    stats_table.autofit = False
    
    stats_data = [
        ("On-Time Delivery", "3-MO. SCOPE IN 2-3 WEEKS (AGILE)"),
        ("ERP & Digital Transform", "900+ MODULES & 16+ UPGRADES"),
        ("Complex Integrations", "SAP, SFA, 7 GATEWAYS (1-MO. SPRINTS)"),
        ("DevOps & Infrastructure", "300+ LIVE DEPLOYMENTS (CI/CD)")
    ]
    
    col_w = Inches(1.81)
    for i, (val, label) in enumerate(stats_data):
        c = stats_table.cell(0, i)
        c.width = col_w
        set_cell_background(c, "F1F5F9")
        clear_cell_borders(c)
        set_cell_margins(c, top=35, bottom=35, left=25, right=25)
        
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0.5)
        rv = p.add_run(val)
        rv.font.name = "Arial"
        rv.font.size = Pt(8.6)
        rv.font.bold = True
        rv.font.color.rgb = NAVY

        p2 = c.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(0)
        rl = p2.add_run(label)
        rl.font.name = "Calibri"
        rl.font.size = Pt(6.3)
        rl.font.bold = True
        rl.font.color.rgb = BLUE

    # Section Heading Helper
    def add_section_heading(title_text, space_before=4.5):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        
        bdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CBD5E1"/></w:pBdr>')
        p._p.get_or_add_pPr().append(bdr)

        r = p.add_run(title_text.upper())
        r.font.name = "Arial"
        r.font.size = Pt(9.0)
        r.font.bold = True
        r.font.color.rgb = NAVY

    # Bullet helper
    def add_bullet(bold_prefix="", body_text="", space_after=1.0):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.09
        
        if bold_prefix:
            rb = p.add_run(bold_prefix + " ")
            rb.font.name = "Calibri"
            rb.font.size = Pt(8.4)
            rb.font.bold = True
            rb.font.color.rgb = DARK
        
        if body_text:
            rt = p.add_run(body_text)
            rt.font.name = "Calibri"
            rt.font.size = Pt(8.4)
            rt.font.color.rgb = DARK
        return p

    # 3. EXECUTIVE SUMMARY
    add_section_heading("Executive Summary", space_before=4)
    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.space_before = Pt(1)
    p_sum.paragraph_format.space_after = Pt(2.5)
    p_sum.paragraph_format.line_spacing = 1.11
    p_sum.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    r_sum = p_sum.add_run(
        "Results-driven Senior Project Manager and Technology Lead with 9+ years of professional experience directing "
        "software engineering teams, enterprise digital transformations, and full-lifecycle software delivery. Proven track record "
        "leading cross-functional teams of software engineers, Odoo developers, business analysts, QA, and DevOps specialists to deliver "
        "secure, scalable, and on-time technology solutions. Certified Odoo Consultant with deep expertise in ERP implementation lifecycles, "
        "business process modeling, system architecture, API/third-party platform integrations (SAP, mobile SFA, payment gateways), "
        "and cloud infrastructure. Established high-efficiency delivery governance compressing 3-month project scopes into 2–3 weeks "
        "with near-100% first-pass UAT acceptance via the Agentic E3 Framework. Adept at translating organizational business objectives "
        "into technical roadmaps, identifying and mitigating delivery risks, and managing stakeholder expectations across 85+ global "
        "and domestic enterprise projects."
    )
    r_sum.font.name = "Calibri"
    r_sum.font.size = Pt(8.4)
    r_sum.font.color.rgb = DARK

    # 4. CORE TECHNICAL & PROJECT LEADERSHIP COMPETENCIES
    add_section_heading("Core Technical & Project Leadership Competencies", space_before=3.5)

    comp_table = doc.add_table(rows=3, cols=2)
    comp_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    comp_table.autofit = False

    comp_items = [
        ("Project & Delivery Management:", "Agile/Scrum, sprint planning, 6-milestone delivery governance, risk management & mitigation, resource allocation, UAT sign-off, SDLC management, on-time delivery KPIs."),
        ("Software Architecture & Code Quality:", "System architecture, modular design patterns, code reviews, coding standards, database design (PostgreSQL/MySQL), API design (REST, JSON-RPC, XML-RPC, Webhooks)."),
        ("ERP & Odoo Technology Leadership:", "End-to-end ERP implementation lifecycle, 900+ custom modules, Odoo Studio, configuration, multi-version database upgrades (v13 to v19) with historical data retention."),
        ("Enterprise System Integrations:", "Bidirectional enterprise connectors (SAP ERP, third-party mobile SFA, WMS, legacy databases), international payment gateways (7 NZ/AUS gateways in 1 month), message queues."),
        ("Infrastructure, Security & DevOps:", "Docker containerization, CI/CD pipelines, Bash automation, Linux/Windows Server (AD/DNS/IIS), Nginx worker isolation, access control, disaster recovery, data protection."),
        ("Team Mentorship & Governance:", "Technical hiring & evaluation, engineer mentorship, performance appraisals, technical documentation, version control (Git), release and change management.")
    ]

    for idx, (head, desc) in enumerate(comp_items):
        row = idx // 2
        col = idx % 2
        cell = comp_table.cell(row, col)
        cell.width = Inches(3.63)
        clear_cell_borders(cell)
        set_cell_margins(cell, top=12, bottom=16, left=20, right=30)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0.5)
        p.paragraph_format.line_spacing = 1.08
        
        rh = p.add_run("• " + head + " ")
        rh.font.name = "Calibri"
        rh.font.size = Pt(8.2)
        rh.font.bold = True
        rh.font.color.rgb = BLUE

        rd = p.add_run(desc)
        rd.font.name = "Calibri"
        rd.font.size = Pt(8.2)
        rd.font.color.rgb = DARK

    # 5. KEY PROJECT DELIVERY & TECHNICAL MILESTONES
    add_section_heading("Key Project Delivery & Technical Milestones", space_before=3.5)

    add_bullet("Enterprise SAP Bidirectional Integration (Delivered in 1 Month):", 
               "Managed and led the technical delivery of an end-to-end integration between Odoo ERP and SAP enterprise software. Synchronized customer masters, sales orders, delivery fulfillments, and payments with automated error handling and logging within a 1-month milestone.")

    add_bullet("7 NZ & AUS Payment Gateways (1-Month Delivery Sprint):", 
               "Directed the development, integration, and certification of 7 distinct New Zealand and Australian payment gateways and BNPL platforms within a single 1-month milestone. Automated invoice reconciliation, webhooks, and PCI-DSS compliance.")

    add_bullet("Hyper-Velocity Delivery Framework (3-Month Scope in 2–3 Weeks):", 
               "Established a structured project governance framework combining modular architecture, automated test suites, and local AI agent workflows, cutting project delivery cycles by 75% while maintaining near-zero defect escape rates.")

    add_bullet("Historical Data & Schema Retention in ERP Migrations:", 
               "Led 16+ major multi-version database upgrades (Odoo 13 to 19). Developed an AST-based schema extraction strategy to retain all historical custom fields, financial ledger records, and audit reports while retiring legacy technical debt.")

    add_bullet("Third-Party Mobile SFA Integration for FMCG Manufacturing:", 
               "Led the integration of an external mobile Sales Force Automation system with Odoo Enterprise for a leading FMCG manufacturer. Managed real-time bulk order fetching, stock reservation, and automated delivery order creation.")

    add_bullet("Infrastructure Automation for 300+ Live Deployments:", 
               "Designed Docker and Bash deployment pipelines managing 300+ live database instances, reducing deployment onboarding time by 85% and ensuring high availability.")

    add_bullet("Multichannel E-Commerce & Marketplace Platforms:", 
               "Managed real-time API sync connecting ERP with international e-commerce channels (Kogan, Tradevine). Automated catalog mapping, inventory sync, and order fulfillment webhooks.")

    add_bullet("Universal Odoo Database Model Context Protocol (MCP) Server:", 
               "Led development of an enterprise MCP server exposing database schemas and models to AI tools for automated schema inspection, task execution, and system modernization.")

    # Explicit page break to guarantee pristine 2-page split
    doc.add_page_break()

    # 6. PROFESSIONAL EXPERIENCE & PROJECT LEADERSHIP HISTORY (PAGE 2)
    add_section_heading("Professional Experience & Project Leadership History", space_before=0)

    # Role 1
    p_r1 = doc.add_paragraph()
    p_r1.paragraph_format.space_before = Pt(2.5)
    p_r1.paragraph_format.space_after = Pt(1)
    p_r1.paragraph_format.keep_with_next = True
    
    r_role = p_r1.add_run("Head of Technology & Technical Project Lead  |  Metamorphosis Ltd.")
    r_role.font.name = "Arial"
    r_role.font.size = Pt(9.6)
    r_role.font.bold = True
    r_role.font.color.rgb = NAVY
    
    r_date = p_r1.add_run("  [Feb 2026 – Present]")
    r_date.font.name = "Calibri"
    r_date.font.size = Pt(8.6)
    r_date.font.bold = True
    r_date.font.color.rgb = BLUE

    add_bullet("Technology Direction & Delivery Management:", 
               "Directs software engineering teams, overseeing the design, development, testing, deployment, and maintenance of enterprise technology solutions for 85+ global and domestic clients.")

    add_bullet("Project Governance & On-Time Delivery:", 
               "Translates business requirements into technical architectures and practical project plans. Monitors timelines, resource allocation, dependencies, and deliverables to ensure in-time project delivery.")

    add_bullet("Architecture & Code Quality Assurance:", 
               "Conducts system architecture reviews, database designs, API specifications, and code reviews. Enforces coding standards, security guidelines, and the Agentic E3 Framework to minimize technical debt.")

    add_bullet("High-Impact Integration Delivery:", 
               "Managed delivery of critical integrations on aggressive timelines, including an end-to-end SAP ERP sync (orders, customers, deliveries, payments) in 1 month, and 7 NZ/AUS payment gateways in a 1-month milestone.")

    add_bullet("Team Mentorship & Capability Building:", 
               "Mentors software engineers, functional consultants, and QA staff. Evaluates technical performance, leads knowledge-sharing sessions, and directs technical recruitment.")

    add_bullet("Infrastructure & Worker Isolation:", 
               "Configured Nginx reverse proxy worker isolation to separate heavy external API sync and e-commerce traffic from internal administrative back-office Odoo workers, eliminating system downtime.")

    # Role 2
    p_r2 = doc.add_paragraph()
    p_r2.paragraph_format.space_before = Pt(4)
    p_r2.paragraph_format.space_after = Pt(1)
    p_r2.paragraph_format.keep_with_next = True
    
    r_role2 = p_r2.add_run("Sr. Systems Engineer & DevOps Lead  |  Metamorphosis Ltd.")
    r_role2.font.name = "Arial"
    r_role2.font.size = Pt(9.2)
    r_role2.font.bold = True
    r_role2.font.color.rgb = NAVY
    
    r_date2 = p_r2.add_run("  [Oct 2023 – Jan 2026]")
    r_date2.font.name = "Calibri"
    r_date2.font.size = Pt(8.6)
    r_date2.font.bold = True
    r_date2.font.color.rgb = BLUE

    add_bullet("Infrastructure Planning & DevOps Automation:", 
               "Designed automated Docker and Bash deployment pipelines managing 300+ live instances, reducing client onboarding time by 85% and ensuring high system stability and reliability.")

    add_bullet("Marketplace & Enterprise SFA Integrations:", 
               "Led real-time sync modules connecting ERP to international e-commerce channels (Kogan, Tradevine, BNPL) and mobile SFA backends. Built automated reconciliation scripts to prevent data drift.")

    add_bullet("Technical Standards & Release Management:", 
               "Maintained version control (Git), release management, change management, and automated testing pipelines across production environments.")

    # Role 3
    p_r3 = doc.add_paragraph()
    p_r3.paragraph_format.space_before = Pt(4)
    p_r3.paragraph_format.space_after = Pt(1)
    p_r3.paragraph_format.keep_with_next = True
    
    r_role3 = p_r3.add_run("Full Stack Developer & Module Lead  |  Metamorphosis Ltd.")
    r_role3.font.name = "Arial"
    r_role3.font.size = Pt(9.2)
    r_role3.font.bold = True
    r_role3.font.color.rgb = NAVY
    
    r_date3 = p_r3.add_run("  [2022 – Sep 2023]")
    r_date3.font.name = "Calibri"
    r_date3.font.size = Pt(8.6)
    r_date3.font.bold = True
    r_date3.font.color.rgb = BLUE

    add_bullet("Business Requirement Translation:", 
               "Hand-coded hundreds of custom Odoo modules matching client operational requirements, business process automations, dynamic reporting, and portal features.")

    add_bullet("Payment Gateways & Database Optimization:", 
               "Built secure payment gateway integrations (cards, mobile banking) and optimized complex multi-table PostgreSQL queries and indexing for high-volume transactions.")

    # Role 4 & 5
    p_r4 = doc.add_paragraph()
    p_r4.paragraph_format.space_before = Pt(4)
    p_r4.paragraph_format.space_after = Pt(1)
    p_r4.paragraph_format.keep_with_next = True
    
    r_role4 = p_r4.add_run("Software Engineer & Systems Specialist  |  MetaKave")
    r_role4.font.name = "Arial"
    r_role4.font.size = Pt(9.2)
    r_role4.font.bold = True
    r_role4.font.color.rgb = NAVY
    
    r_date4 = p_r4.add_run("  [2020 – 2021]")
    r_date4.font.name = "Calibri"
    r_date4.font.size = Pt(8.6)
    r_date4.font.bold = True
    r_date4.font.color.rgb = BLUE

    add_bullet("Application Development & Mobile APIs:", 
               "Engineered backend web services and REST endpoints for Magento 2+ e-commerce platforms connected to mobile apps. Managed cloud server environments and Nginx SSL reverse proxies.")

    p_r5 = doc.add_paragraph()
    p_r5.paragraph_format.space_before = Pt(4)
    p_r5.paragraph_format.space_after = Pt(1)
    p_r5.paragraph_format.keep_with_next = True
    
    r_role5 = p_r5.add_run("IT Executive  |  Aptech Industries")
    r_role5.font.name = "Arial"
    r_role5.font.size = Pt(9.2)
    r_role5.font.bold = True
    r_role5.font.color.rgb = NAVY
    
    r_date5 = p_r5.add_run("  [Jun 2017 – Sep 2019]")
    r_date5.font.name = "Calibri"
    r_date5.font.size = Pt(8.6)
    r_date5.font.bold = True
    r_date5.font.color.rgb = BLUE

    add_bullet("Enterprise IT & Server Administration:", 
               "Managed enterprise IT operations, Windows Server administration (Active Directory, DNS, IIS, Group Policy up to 2019), network security, disaster recovery, and infrastructure deployments.")

    # 7. EDUCATION, CERTIFICATIONS & RESEARCH
    add_section_heading("Education, Certifications & Research", space_before=4)

    add_bullet("Certified Odoo Consultant:", "Metamorphosis Ltd. (Official Odoo Silver Partner). Extensive expertise in Odoo standard architecture, ORM, Studio, configuration, and ERP implementation lifecycles.")
    add_bullet("Master of Information Technology (MIT):", "Jahangirnagar University (2018 – 2019). Focus on Enterprise Software Systems, Advanced Database Management, and Distributed Computing.")
    add_bullet("Bachelor of Science (B.Sc):", "Shahjalal University of Science and Technology (SUST). 1st Major in Mathematics, 2nd Major in Computer Science & Engineering (CSE).")
    add_bullet("IEEE Research Author (2021):", "Published paper on data processing optimization in Hadoop MapReduce, introducing the IMC algorithm (65% efficiency gain). Applied distributed principles to enterprise database scale.")

    # Footer note
    p_foot = doc.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_foot.paragraph_format.space_before = Pt(4)
    p_foot.paragraph_format.space_after = Pt(0)
    rf = p_foot.add_run("Md Niaj Shahriar Shishir  |  Senior Project Manager & Technology Lead  |  Portfolio: nsrshishir.github.io/nsrshishir")
    rf.font.name = "Calibri"
    rf.font.size = Pt(7.0)
    rf.font.color.rgb = MUTED

    doc.save(output_path)
    print(f"Successfully generated SPM DOCX: {output_path}")

if __name__ == "__main__":
    output_docx = "/home/niaj/Documents/nsrshishir/Md_Niaj_Shahriar_Shishir_Senior_Project_Manager.docx"
    create_spm_docx(output_docx)
